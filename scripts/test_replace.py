from docx import Document
import re
import zipfile

template_path = 'backend/templates/word/attestation_presence_template.docx'

def simulate():
    nom_complet = 'DUPONT'
    poste = 'Technicien'
    date_prise_service = '02 février 2015'
    date_aujourdhui = '01/06/2026'
    annee = '2026'
    ref_number = f"{annee}0601181246"
    reference = f"{ref_number}/MND/DPAF/SRHDS/SA"
    placeholders = {
        '{{NOM_COMPLET}}': nom_complet,
        '{{POSTE}}': poste,
        '{{DATE_PRISE_SERVICE}}': date_prise_service,
        '{{DATE_AUJOURD_HUI}}': date_aujourdhui,
        '_AUJOURD_HUI}}': date_aujourdhui,
    }

    doc = Document(template_path)
    found = []
    for paragraph in doc.paragraphs:
        para_text = ''.join([r.text for r in paragraph.runs])
        new_text = para_text
        if '{{REFERENCE}}' in new_text:
            if '/MND/DPAF/SRHDS/SA' in new_text:
                ref_display = ref_number
            else:
                ref_display = reference
            new_text = new_text.replace('{{REFERENCE}}', ref_display)
        for key, val in placeholders.items():
            if key in new_text:
                new_text = new_text.replace(key, val)
        new_text = re.sub(r"(\d{2}/\d{2}/\d{4})\1+", r"\1", new_text)
        if new_text != para_text and (date_aujourdhui in new_text or ref_number in new_text):
            found.append((para_text, new_text))
    for before, after in found:
        print('--- BEFORE:')
        print(before)
        print('\n--- AFTER:')
        print(after)
        print('\n')

if __name__ == '__main__':
    simulate()
